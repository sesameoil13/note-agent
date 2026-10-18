"""Independent Word fixtures authored by python-docx (not the Dart extractor).

Run with bundled Python / python-docx 1.2.0. No network or private content.
"""
from pathlib import Path
from docx import Document

ROOT = Path(__file__).resolve().parents[1] / 'test' / 'fixtures' / 'word'
ROOT.mkdir(parents=True, exist_ok=True)
doc = Document()
doc.add_paragraph('中文资料：随口记录想法。 Latin 42 😀')
table = doc.add_table(rows=3, cols=3)
for row, values in zip(table.rows, [['产品', '数量', '备注'], ['苹果', '2', 'A | B'], ['茶叶', '', '第一行\n第二行']]):
    for cell, value in zip(row.cells, values): cell.text = value
doc.add_paragraph('表格后的原文。')
doc.save(ROOT / 'body-table.docx')

# Known tiny PNG authored locally; no screenshot/user data.
from base64 import b64decode
from io import BytesIO
png = b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+jUZkAAAAASUVORK5CYII=')
doc = Document()
doc.add_paragraph('配图前的文字。')
doc.add_picture(BytesIO(png))
doc.add_paragraph('配图后的文字。')
doc.sections[0].header.paragraphs[0].text = '页眉中的额外内容'
doc.save(ROOT / 'mixed.docx')
doc = Document()
doc.add_picture(BytesIO(png))
doc.save(ROOT / 'images-only.docx')
doc = Document()
doc.add_paragraph('复杂表格前。')
t = doc.add_table(rows=2, cols=2)
t.cell(0, 0).merge(t.cell(0, 1)).text = '跨列内容'
t.cell(1, 0).text = '外层单元格'
nested = t.cell(1, 1).add_table(rows=1, cols=1)
nested.cell(0, 0).text = '嵌套内容'
doc.add_paragraph('复杂表格后。')
doc.save(ROOT / 'complex.docx')

from zipfile import ZipFile, ZIP_DEFLATED
W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
def variant(name, body, *, main='word/document.xml', namespace=W):
    with ZipFile(ROOT / 'body-table.docx') as z:
        entries = {n: z.read(n) for n in z.namelist()}
    entries['word/document.xml'] = f'<?xml version="1.0"?><x:document xmlns:x="{namespace}"><x:body>{body}</x:body></x:document>'.encode()
    if main != 'word/document.xml':
        entries[main] = entries.pop('word/document.xml')
        entries['_rels/.rels'] = entries['_rels/.rels'].replace(b'word/document.xml', main.encode())
        entries['[Content_Types].xml'] = entries['[Content_Types].xml'].replace(b'word/document.xml', main.encode())
        entries.pop('word/_rels/document.xml.rels', None)
    with ZipFile(ROOT / name, 'w', compression=ZIP_DEFLATED) as z:
        for n, data in entries.items(): z.writestr(n, data)

variant('revisions.docx', '''<x:p><x:r><x:t>当前段落</x:t></x:r><x:ins><x:r><x:t>新增文字</x:t></x:r></x:ins><x:del><x:r><x:delText>已删除文字</x:delText></x:r></x:del><x:r><x:fldChar x:fldCharType="begin"/><x:instrText>DATE</x:instrText><x:fldChar x:fldCharType="separate"/><x:t>字段缓存</x:t><x:fldChar x:fldCharType="end"/></x:r></x:p><x:sdt><x:sdtContent><x:p><x:r><x:t>控件文字</x:t></x:r></x:p></x:sdtContent></x:sdt><x:altChunk xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" r:id="unread"/>''')
variant('strict-renamed.docx', '<x:p><x:r><x:t>严格命名空间和非默认主路径。</x:t></x:r></x:p>', main='custom/main.xml', namespace='http://purl.oclc.org/ooxml/wordprocessingml/main')

# Real packaged corruption/resource probes, all containing synthetic content.
from struct import unpack_from, pack_into
with ZipFile(ROOT / 'body-table.docx') as z:
    base = {n: z.read(n) for n in z.namelist()}
def package(name, entries, **kwargs):
    with ZipFile(ROOT / name, 'w', compression=ZIP_DEFLATED) as z:
        for n, data in entries: z.writestr(n, data)
def changed_main(name, xml):
    entries = dict(base); entries['word/document.xml'] = xml
    package(name, entries.items())
changed_main('damaged-xml.docx', b'<document><not-closed>')
changed_main('undeclared-prefix.docx', b'<x:document><x:body/></x:document>')
changed_main('dtd.docx', b'<!DOCTYPE x:document [<!ENTITY remote SYSTEM "https://invalid.example/never-load">]><x:document xmlns:x="'+W.encode()+b'"><x:body><x:p><x:r><x:t>&remote;</x:t></x:r></x:p></x:body></x:document>')
changed_main('deep.docx', (f'<x:document xmlns:x="{W}"><x:body>'+ '<x:sdt>'*70 + '<x:p><x:r><x:t>深层</x:t></x:r></x:p>' + '</x:sdt>'*70 + '</x:body></x:document>').encode())
changed_main('expanded-limit.docx', b' ' * (4 * 1024 * 1024 + 1))
package('duplicate.docx', list(base.items()) + [('word/document.xml', base['word/document.xml'])])
package('traversal.docx', list(base.items()) + [('../escape.txt', b'no escape')])
package('missing-main.docx', ((n, d) for n, d in base.items() if n != 'word/document.xml'))
package('entry-limit.docx', list(base.items()) + [(f'empty/{i}', b'') for i in range(2050)])
(ROOT / 'broken.docx').write_bytes(b'PK\x03\x04broken package')
(ROOT / 'legacy.doc').write_bytes(bytes.fromhex('d0cf11e0a1b11ae1') + bytes(504))
for name, kind in [('crc.docx', 'crc'), ('encrypted.docx', 'encrypted'), ('lying-size.docx', 'size')]:
    raw = bytearray((ROOT / 'body-table.docx').read_bytes())
    end = raw.rfind(b'PK\x05\x06'); offset = unpack_from('<I', raw, end+16)[0]
    while raw[offset:offset+4] == b'PK\x01\x02':
        nl, el, cl = unpack_from('<HHH', raw, offset+28)
        if raw[offset+46:offset+46+nl] == b'word/document.xml':
            local = unpack_from('<I', raw, offset+42)[0]
            if kind == 'crc':
                value = unpack_from('<I', raw, offset+16)[0] ^ 0x12345678
                pack_into('<I', raw, offset+16, value); pack_into('<I', raw, local+14, value)
            elif kind == 'encrypted':
                flags = unpack_from('<H', raw, offset+8)[0] | 1
                pack_into('<H', raw, offset+8, flags); pack_into('<H', raw, local+6, flags)
            else:
                pack_into('<I', raw, offset+24, 1); pack_into('<I', raw, local+22, 1)
            break
        offset += 46+nl+el+cl
    (ROOT / name).write_bytes(raw)
variant('long.docx', ''.join(f'<x:p><x:r><x:t>第{i}段原文。</x:t></x:r></x:p>' for i in range(30000)))

# Independent read-back: python-docx can reopen the valid documents and confirms
# literal cells, including newline/empty cells; damaged fixtures are intentional.
assert Document(ROOT / 'body-table.docx').tables[0].cell(2, 2).text == '第一行\n第二行'
assert Document(ROOT / 'mixed.docx').sections[0].header.paragraphs[0].text == '页眉中的额外内容'
assert Document(ROOT / 'complex.docx').tables[0].cell(1, 1).tables[0].cell(0, 0).text == '嵌套内容'

# ZIP data descriptors are used by streaming Office exporters.
class StreamingBuffer(BytesIO):
    def seek(self, *args): raise OSError('streaming fixture')
    def seekable(self): return False
stream = StreamingBuffer()
with ZipFile(stream, 'w', compression=ZIP_DEFLATED) as z:
    for n, data in base.items(): z.writestr(n, data)
(ROOT / 'streamed.docx').write_bytes(stream.getvalue())
raw = bytearray(stream.getvalue())
end = raw.rfind(b'PK\x05\x06'); offset = unpack_from('<I', raw, end+16)[0]
while raw[offset:offset+4] == b'PK\x01\x02':
    nl, el, cl = unpack_from('<HHH', raw, offset+28)
    if raw[offset+46:offset+46+nl] == b'word/document.xml':
        local = unpack_from('<I', raw, offset+42)[0]
        ln, le = unpack_from('<HH', raw, local+26)
        descriptor = local + 30 + ln + le + unpack_from('<I', raw, offset+20)[0]
        assert raw[descriptor:descriptor+4] == b'PK\x07\x08'
        raw[descriptor+4] ^= 1
        break
    offset += 46+nl+el+cl
(ROOT / 'bad-descriptor.docx').write_bytes(raw)
changed_main('utf16.docx', (f'<?xml version="1.0" encoding="UTF-16"?><x:document xmlns:x="{W}"><x:body><x:p><x:r><x:t>UTF16 中文来源 😀</x:t></x:r></x:p></x:body></x:document>').encode('utf-16'))
variant('alternate.docx', '''<x:p><x:r><x:t>替代内容：</x:t></x:r><mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" xmlns:feature="urn:fixture-feature"><mc:Choice Requires="feature"><x:r><x:t>主分支文字</x:t></x:r></mc:Choice><mc:Fallback><x:r><x:t>回退分支文字</x:t></x:r></mc:Fallback></mc:AlternateContent></x:p><x:p><x:pPr><x:numPr><x:ilvl x:val="0"/><x:numId x:val="1"/></x:numPr></x:pPr><x:r><x:t>列表正文</x:t><x:noBreakHyphen/><x:t>续文</x:t></x:r></x:p>''')
package('case-duplicate.docx', list(base.items()) + [('WORD/DOCUMENT.XML', base['word/document.xml'])])
cased = dict(base)
cased['_rels/.rels'] = cased['_rels/.rels'].replace(b'word/document.xml', b'WORD/DOCUMENT.XML')
cased['[Content_Types].xml'] = cased['[Content_Types].xml'].replace(b'word/document.xml', b'Word/Document.Xml')
package('case-reference.docx', cased.items())

# Merger review regressions: legal content controls and omitted grid positions
# are not equivalent to a simple rectangular table. These author the same
# independently observed structures without using the production reader.
from xml.etree import ElementTree as ET
word_tag = lambda name: '{' + W + '}' + name
controlled = ET.fromstring(base['word/document.xml'])
cell = controlled.find('.//' + word_tag('tc'))
paragraph = cell.find(word_tag('p'))
cell.remove(paragraph)
control = ET.SubElement(cell, word_tag('sdt'))
ET.SubElement(control, word_tag('sdtContent')).append(paragraph)
changed_main('cell-content-control.docx', ET.tostring(controlled, encoding='utf-8', xml_declaration=True))
omitted = ET.fromstring(base['word/document.xml'])
table = omitted.find('.//' + word_tag('tbl'))
grid = table.find(word_tag('tblGrid'))
ET.SubElement(grid, word_tag('gridCol'), {word_tag('w'): '2000'})
for index, row in enumerate(table.findall(word_tag('tr'))):
    props = row.find(word_tag('trPr'))
    if props is None:
        props = ET.Element(word_tag('trPr')); row.insert(0, props)
    ET.SubElement(props, word_tag('gridBefore' if index == 0 else 'gridAfter'), {word_tag('val'): '1'})
changed_main('omitted-grid-cells.docx', ET.tostring(omitted, encoding='utf-8', xml_declaration=True))
